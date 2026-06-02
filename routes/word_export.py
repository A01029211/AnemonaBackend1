"""
word_export.py  —  Generador de .docx para el SRS dinámico de Anémona/Banorte
"""

import io
import os
from typing import Any

import matplotlib
matplotlib.use("Agg")   # backend sin pantalla, obligatorio en servidor
import matplotlib.pyplot as plt

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT


AZUL_BANORTE = RGBColor(0x13, 0x3B, 0x73)
ROJO_BANORTE = RGBColor(0xEB, 0x00, 0x29)
AZUL_TEXTO   = RGBColor(0x1D, 0x5D, 0xA8)
GRIS_LABEL   = RGBColor(0x7C, 0x7C, 0x7C)
BLANCO       = RGBColor(0xFF, 0xFF, 0xFF)


# ─── Normalización ───────────────────────────────────────────────────────────

def _val(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, list):
        return ", ".join(_val(i) for i in v)
    s = str(v).strip()
    return "" if s.upper() == "NULL" else s


# ─── Helpers de estilo ───────────────────────────────────────────────────────

def _font(run, size_pt: float, bold=False, italic=False, color: RGBColor = None):
    run.font.name = "Arial"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def _para(container, text="", size=11, bold=False, italic=False,
          color=None, before=0, after=6, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after  = Pt(after)
    if text:
        r = p.add_run(text)
        _font(r, size, bold=bold, italic=italic, color=color)
    return p


def _cell_bg(cell, hex6: str):
    tc   = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd  = OxmlElement("w:shd")
    shd.set(qn("w:val"),   "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"),  hex6)
    tcPr.append(shd)


def _border_top(p, color_hex="000000", size=12):
    pPr  = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top  = OxmlElement("w:top")
    top.set(qn("w:val"),   "single")
    top.set(qn("w:sz"),    str(size))
    top.set(qn("w:space"), "4")
    top.set(qn("w:color"), color_hex)
    pBdr.append(top)
    pPr.append(pBdr)


def _spacer(doc):
    _para(doc, after=2)


# ─── Componentes UI ──────────────────────────────────────────────────────────

def _section_line(doc, number: str, title: str, optional=True):
    """Línea con borde superior negro, número + título grande — para W006/WChart"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    _border_top(p)
    r1 = p.add_run(f"{number}  ")
    _font(r1, 13, bold=True)
    r2 = p.add_run(title)
    _font(r2, 13, bold=True)
    if optional:
        r3 = p.add_run("   (Opcional)")
        _font(r3, 9, color=ROJO_BANORTE)


def _subsection_line(doc, number: str, title: str, optional=False):
    """Línea con borde superior, número + título — para W003/W005"""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(4)
    _border_top(p)
    r1 = p.add_run(f"{number}  ")
    _font(r1, 13, bold=True)
    r2 = p.add_run(f"{title}.")
    _font(r2, 13, bold=True)
    if optional:
        r3 = p.add_run("   (Opcional)")
        _font(r3, 9, color=ROJO_BANORTE)


# ─── Intro ───────────────────────────────────────────────────────────────────

INTRO = (
    "Este cuestionario tiene como propósito conocer cuáles son los beneficios, costos y riesgos "
    "relacionados con cada iniciativa que ingresa al portafolio de proyectos y mantenimientos "
    "tecnológicos de Áreas de Soporte. Esta información será de utilidad para ponderar el "
    "portafolio en su conjunto y priorizar la atención de los requerimientos de acuerdo a su "
    "beneficio económico, alineación estratégica y conveniencia de su realización."
)

def _add_intro(doc):
    _para(doc, INTRO, size=11, after=12)


# ─── W000 ────────────────────────────────────────────────────────────────────

W000_BASE = [
    "SOLICITANTE", "INFO_CONTACTO", "DGA", "PATROCINADOR",
    "CR", "SOCIO", "NOMBRE_INICIATIVA", "TIPO_INICIATIVA",
]

def _render_w000(doc, widget: dict):
    campos = widget.get("campos", {}) or {}
    desc   = widget.get("descripcion_campos", {}) or {}

    base_keys  = [k for k in W000_BASE if k in campos]
    extra_keys = [k for k in campos if k not in W000_BASE]
    all_keys   = base_keys + extra_keys
    if not all_keys:
        return

    tbl = doc.add_table(rows=len(all_keys), cols=2)
    tbl.style     = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    for i, key in enumerate(all_keys):
        row   = tbl.rows[i]
        label = desc.get(key) or key.replace("_", " ").title()
        valor = _val(campos.get(key))

        row.cells[0].width = Inches(2.6)
        row.cells[1].width = Inches(4.9)

        r0 = row.cells[0].paragraphs[0].add_run(label)
        _font(r0, 11, bold=True)

        r1 = row.cells[1].paragraphs[0].add_run(valor)
        _font(r1, 11)

    _spacer(doc)


# ─── W003 ────────────────────────────────────────────────────────────────────

def _normalizar_headers(raw: Any) -> list:
    """Acepta headers en cualquier formato que mande Firestore."""
    if not raw:
        return []
    if isinstance(raw, dict):
        return [{"key": k, "label": v} for k, v in raw.items()]
    if isinstance(raw, list):
        normalized = []
        for item in raw:
            if isinstance(item, dict):
                key   = str(item.get("key",   item.get("Key",   "")))
                label = str(item.get("label", item.get("Label", key)))
                normalized.append({"key": key, "label": label})
            else:
                s = str(item).strip()
                normalized.append({"key": s, "label": s})
        return normalized
    return []


def _render_w003(doc, widget: dict):
    campos  = widget.get("campos", {}) or {}
    pos     = widget.get("posicion", "")
    titulo  = _val(campos.get("titulo")) or widget.get("titulo", "")
    filas   = campos.get("filas") or []

    # ── Resolver headers ────────────────────────────────────────────────────
    headers = _normalizar_headers(campos.get("headers"))

    # Si no vienen headers explícitos, inferirlos desde la primera fila
    if not headers and filas:
        primera = filas[0]
        if isinstance(primera, dict):
            headers = [{"key": k, "label": k} for k in primera.keys()]

    # Fallback final
    if not headers:
        headers = [
            {"key": "TIPO",             "label": "Riesgo"},
            {"key": "PROBABLE_PERDIDA", "label": "Probable Pérdida"},
            {"key": "JUSTIFICACION",    "label": "Justificación"},
        ]

    # Filtrar columnas fantasma (key vacío que el frontend a veces manda)
    headers = [h for h in headers if h.get("key", "").strip()]

    _subsection_line(doc, str(pos), titulo, optional=True)

    num_cols = len(headers)
    tbl = doc.add_table(rows=1 + max(len(filas), 1), cols=num_cols)
    tbl.style     = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Header azul
    hrow = tbl.rows[0]
    for i, h in enumerate(headers):
        cell = hrow.cells[i]
        _cell_bg(cell, "133b73")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h["label"])
        _font(r, 11, bold=True, color=BLANCO)

    if filas:
        for ri, fila in enumerate(filas):
            row = tbl.rows[ri + 1]
            if not isinstance(fila, dict):
                continue
            for ci, h in enumerate(headers):
                valor = _val(fila.get(h["key"], ""))
                r = row.cells[ci].paragraphs[0].add_run(valor)
                _font(r, 11)
    else:
        merged = tbl.rows[1].cells[0].merge(tbl.rows[1].cells[-1])
        p = merged.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("N/A")
        _font(r, 11, color=GRIS_LABEL)

    _spacer(doc)


# ─── W005 ────────────────────────────────────────────────────────────────────

def _render_w005(doc, widget: dict):
    campos = widget.get("campos", {}) or {}
    pos    = widget.get("posicion", "")
    titulo = _val(campos.get("titulo")) or widget.get("titulo", "")
    filas  = campos.get("filas") or []

    _subsection_line(doc, str(pos), titulo)

    if not filas:
        _spacer(doc)
        return

    # Calcular el máximo de celdas reales en cualquier fila
    max_celdas = max((len(f.get("celdas", [])) for f in filas), default=1)
    # Siempre usar max_celdas columnas (mínimo 1)
    num_cols = max(max_celdas, 1)

    tbl = doc.add_table(rows=len(filas), cols=num_cols)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    for ri, fila in enumerate(filas):
        celdas = fila.get("celdas", [])
        row    = tbl.rows[ri]
        n      = len(celdas)

        if n == 0:
            continue

        # Si hay menos celdas que columnas, mergear el resto en la última
        if n == 1:
            # Celda única: mergea todas las columnas
            cell = row.cells[0]
            if num_cols > 1:
                cell = row.cells[0].merge(row.cells[num_cols - 1])
            cel   = celdas[0]
            texto = _val(cel.get("valor", "")) or _val(cel.get("label", ""))
            bold  = bool(cel.get("bold", False))
            r = cell.paragraphs[0].add_run(texto)
            _font(r, 11, bold=bold)

        elif n < num_cols:
            # Distribuir celdas proporcionalmente mergeando columnas extra
            # Estrategia: primeras (n-1) celdas ocupan 1 col cada una,
            # la última ocupa el resto
            for ci, cel in enumerate(celdas[:-1]):
                texto = _val(cel.get("valor", "")) or _val(cel.get("label", ""))
                bold  = bool(cel.get("bold", False))
                r = row.cells[ci].paragraphs[0].add_run(texto)
                _font(r, 11, bold=bold, color=GRIS_LABEL if not bold else None)
            # Última celda: mergea columnas restantes
            cel_last  = celdas[-1]
            texto     = _val(cel_last.get("valor", "")) or _val(cel_last.get("label", ""))
            bold      = bool(cel_last.get("bold", False))
            cell_last = row.cells[n - 1].merge(row.cells[num_cols - 1])
            r = cell_last.paragraphs[0].add_run(texto)
            _font(r, 11, bold=bold, color=GRIS_LABEL if not bold else None)

        else:
            # n == num_cols: una celda por columna
            for ci, cel in enumerate(celdas[:num_cols]):
                texto = _val(cel.get("valor", "")) or _val(cel.get("label", ""))
                bold  = bool(cel.get("bold", False))
                r = row.cells[ci].paragraphs[0].add_run(texto)
                _font(r, 11, bold=bold, color=GRIS_LABEL if not bold else None)

    _spacer(doc)


# ─── W004 / WChart ───────────────────────────────────────────────────────────

_PALETTE = [
    "#133b73",  # azul oscuro banorte
    "#EB0029",  # rojo banorte
    "#4a7fc1",  # azul medio
    "#e8a020",  # ámbar
    "#2e7d32",  # verde
    "#7b1fa2",  # morado
    "#00838f",  # teal
    "#c62828",  # rojo oscuro
]

def _generar_grafica_bytes(barras: list, unidad: str = "", pie: str = "") -> bytes:
    """Genera la gráfica de barras como PNG en memoria, igual que el frontend."""
    etiquetas = [str(b.get("etiqueta", "")) for b in barras]
    valores   = [float(b.get("valor", 0)) for b in barras]
    colores   = [_PALETTE[i % len(_PALETTE)] for i in range(len(barras))]

    fig, ax = plt.subplots(figsize=(6.5, 2.8))
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")

    x    = range(len(etiquetas))
    bars = ax.bar(x, valores, color=colores, width=0.55, zorder=3)

    max_val = max(valores) if valores else 1

    # Valores encima de cada barra (bold, color de la barra)
    for bar, val, color in zip(bars, valores, colores):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height() + max_val * 0.02,
            f"{val:g}{unidad}",
            ha="center", va="bottom",
            fontsize=8, fontweight="bold", color=color,
        )

    # Líneas guía horizontales punteadas (25 / 50 / 75 / 100 %)
    for pct in [0.25, 0.5, 0.75, 1.0]:
        ax.axhline(max_val * pct, color="#e0e0e0", linestyle="--", linewidth=0.7, zorder=1)

    # Etiquetas eje X con el color de cada barra
    ax.set_xticks(list(x))
    ax.set_xticklabels(etiquetas, fontsize=8, fontweight="semibold")
    for tick, color in zip(ax.get_xticklabels(), colores):
        tick.set_color(color)

    # Ejes: solo línea inferior negra
    ax.yaxis.set_visible(False)
    for spine in ["top", "left", "right"]:
        ax.spines[spine].set_visible(False)
    ax.spines["bottom"].set_color("black")
    ax.spines["bottom"].set_linewidth(1.5)

    ax.set_xlim(-0.5, len(etiquetas) - 0.5)
    ax.set_ylim(0, max_val * 1.22)

    if pie:
        fig.text(0.01, -0.06, pie, fontsize=7, color="#888888", style="italic")

    plt.tight_layout(pad=0.3)

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def _render_wchart(doc, widget: dict):
    campos = widget.get("campos", {}) or {}
    pos    = widget.get("posicion", "")
    titulo = _val(campos.get("titulo")) or widget.get("titulo", "Gráfica")
    pie    = _val(campos.get("pie", ""))
    unidad = _val(campos.get("unidad", ""))
    barras = campos.get("barras") or []

    _section_line(doc, f"{pos}.", titulo, optional=False)

    if not barras:
        _para(doc, "Sin datos de gráfica.", size=11, italic=True, color=GRIS_LABEL)
        _spacer(doc)
        return

    # Generar imagen y embeberla
    img_bytes = _generar_grafica_bytes(barras, unidad=unidad, pie=pie)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after  = Pt(4)
    run = p.add_run()
    run.add_picture(io.BytesIO(img_bytes), width=Inches(6.0))

    _spacer(doc)


# ─── W006 ────────────────────────────────────────────────────────────────────

def _render_w006(doc, widget: dict):
    campos  = widget.get("campos", {}) or {}
    pos     = widget.get("posicion", "")
    titulo  = _val(campos.get("titulo")) or widget.get("titulo", "")
    bloques = campos.get("bloques") or []

    _section_line(doc, f"{pos}.", titulo, optional=True)

    for bloque in bloques:
        tipo  = _val(bloque.get("tipo", "parrafo"))
        texto = _val(bloque.get("texto", ""))
        if not texto:
            continue

        # Subtítulo → negrita, negro
        if tipo == "subtitulo":
            _para(doc, texto, size=11, bold=True, after=2)
            continue

        # Párrafo: parsear línea por línea para bullets y listas numeradas
        lines = [l for l in texto.split("\n") if l.strip()]
        i = 0
        while i < len(lines):
            line = lines[i].strip()

            # Lista numerada: "1. texto"
            if line and line[0].isdigit() and ". " in line[:6]:
                _, content = line.split(". ", 1)
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after  = Pt(2)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.left_indent  = Inches(0.25)
                r = p.add_run(content)
                _font(r, 11, italic=True, color=AZUL_TEXTO)
                i += 1
                continue

            # Bullet: "- texto"
            if line.startswith("- "):
                content = line[2:]
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after  = Pt(2)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.left_indent  = Inches(0.25)
                r = p.add_run(content)
                _font(r, 11, italic=True, color=AZUL_TEXTO)
                i += 1
                continue

            # Párrafo normal → itálica azul
            _para(doc, line, size=11, italic=True, color=AZUL_TEXTO, after=6)
            i += 1

    _spacer(doc)


# ─── Header / Footer ─────────────────────────────────────────────────────────

def _img_buf_from_b64(b64: str | None) -> io.BytesIO | None:
    """Decodifica base64 → BytesIO listo para add_picture. Devuelve None si falla."""
    if not b64:
        return None
    try:
        import base64
        from PIL import Image as PILImage
        if "," in b64:
            b64 = b64.split(",", 1)[1]
        raw = base64.b64decode(b64)
        img = PILImage.open(io.BytesIO(raw)).convert("RGBA")
        fondo = PILImage.new("RGBA", img.size, (255, 255, 255, 255))
        fondo.paste(img, mask=img.split()[3])
        buf = io.BytesIO()
        fondo.convert("RGB").save(buf, format="PNG")
        buf.seek(0)
        return buf
    except Exception:
        return None


def _img_buf_from_file(path: str) -> io.BytesIO | None:
    """Carga imagen desde disco → BytesIO con fondo blanco. Devuelve None si no existe."""
    if not os.path.exists(path):
        return None
    try:
        from PIL import Image as PILImage
        img = PILImage.open(path).convert("RGBA")
        fondo = PILImage.new("RGBA", img.size, (255, 255, 255, 255))
        fondo.paste(img, mask=img.split()[3])
        buf = io.BytesIO()
        fondo.convert("RGB").save(buf, format="PNG")
        buf.seek(0)
        return buf
    except Exception:
        return None


# Paths relativos a este archivo (routes/word_export.py → ../public/images/)
_BASE_DIR   = os.path.dirname(os.path.abspath(__file__))
_IMG_RAYA   = os.path.join(_BASE_DIR, "..", "public", "images", "rayaNegra.png")
_IMG_FOOTER = os.path.join(_BASE_DIR, "..", "public", "images", "banortegf.png")


def _add_header(doc, img_raya_b64: str | None = None, nombre_plantilla: str = "Levantamiento de Requerimiento"):
    sec    = doc.sections[0]
    header = sec.header
    header.is_linked_to_previous = False

    for p in header.paragraphs:
        p.clear()

    # ── Tabla sin bordes: celda izquierda = texto, celda derecha = imagen ──
    # Ancho total del área de contenido: 8.5" - 1" - 1" = 6.5" = 9360 DXA
    tbl = header.add_table(rows=1, cols=2, width=Inches(6.5))
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT

    cell_left  = tbl.cell(0, 0)
    cell_right = tbl.cell(0, 1)

    # Quitar todos los bordes de ambas celdas
    for cell in [cell_left, cell_right]:
        tc   = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBdr = OxmlElement("w:tcBdr")
        for side in ["top", "left", "bottom", "right", "insideH", "insideV"]:
            border = OxmlElement(f"w:{side}")
            border.set(qn("w:val"), "none")
            border.set(qn("w:sz"), "0")
            border.set(qn("w:space"), "0")
            border.set(qn("w:color"), "auto")
            tcBdr.append(border)
        tcPr.append(tcBdr)

    # Celda izquierda — texto, ocupa el resto del ancho
    img_buf = _img_buf_from_b64(img_raya_b64) or _img_buf_from_file(_IMG_RAYA)
    img_w = Inches(2.2) if img_buf else Inches(0)
    left_w_dxa  = int((6.5 - 2.2) * 1440) if img_buf else 9360
    right_w_dxa = int(2.2 * 1440)          if img_buf else 0

    # Setear ancho de celdas via XML
    for cell, w_dxa in [(cell_left, left_w_dxa), (cell_right, right_w_dxa)]:
        tc   = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcW  = OxmlElement("w:tcW")
        tcW.set(qn("w:w"),    str(w_dxa))
        tcW.set(qn("w:type"), "dxa")
        tcPr.append(tcW)

    p_left = cell_left.paragraphs[0]
    p_left.paragraph_format.space_before = Pt(0)
    p_left.paragraph_format.space_after  = Pt(0)
    r1 = p_left.add_run("Formato Estándar")
    _font(r1, 9, bold=True, color=GRIS_LABEL)
    r2 = p_left.add_run(f"  |  {nombre_plantilla}")
    _font(r2, 9, color=GRIS_LABEL)

    # Celda derecha — imagen alineada a la derecha
    if img_buf:
        p_right = cell_right.paragraphs[0]
        p_right.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_right.paragraph_format.space_before = Pt(0)
        p_right.paragraph_format.space_after  = Pt(0)
        r_img = p_right.add_run()
        r_img.add_picture(img_buf, width=Inches(2.2), height=Inches(0.3))

    # ── Línea inferior ──────────────────────────────────────────────────────
    p_line = header.add_paragraph()
    p_line.paragraph_format.space_before = Pt(3)
    p_line.paragraph_format.space_after  = Pt(0)
    pPr2 = p_line._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bot  = OxmlElement("w:bottom")
    bot.set(qn("w:val"),   "single")
    bot.set(qn("w:sz"),    "6")
    bot.set(qn("w:space"), "4")
    bot.set(qn("w:color"), "CCCCCC")
    pBdr.append(bot)
    pPr2.append(pBdr)


def _add_footer(doc, img_footer_b64: str | None = None):
    sec    = doc.sections[0]
    footer = sec.footer
    footer.is_linked_to_previous = False

    for p in footer.paragraphs:
        p.clear()

    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)

    img_buf = _img_buf_from_b64(img_footer_b64) or _img_buf_from_file(_IMG_FOOTER)
    if img_buf:
        r_img = p.add_run()
        r_img.add_picture(img_buf, width=Inches(3), height=Inches(0.4))
    else:
        r = p.add_run("GRUPO FINANCIERO BANORTE")
        _font(r, 9, bold=True, color=ROJO_BANORTE)


# ─── FastAPI ─────────────────────────────────────────────────────────────────

from fastapi import APIRouter
from fastapi.responses import Response
from pydantic import BaseModel

router = APIRouter()

class ExportarWordRequest(BaseModel):
    doc_id: str
    widgets: list[dict[str, Any]]
    nombre_plantilla: str = "Levantamiento de Requerimiento"
    img_raya_b64:   str | None = None
    img_footer_b64: str | None = None

@router.post("/widgets/exportar-word")
async def exportar_word(body: ExportarWordRequest):
    docx_bytes = generar_word_srs(
        body.widgets,
        nombre_plantilla=body.nombre_plantilla,
        img_raya_b64=body.img_raya_b64,
        img_footer_b64=body.img_footer_b64,
    )
    safe_id = body.doc_id.replace("/", "_").replace(" ", "_")
    return Response(
        content=docx_bytes,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f'attachment; filename="SRS_{safe_id}.docx"',
            "Content-Length": str(len(docx_bytes)),
        },
    )


# ─── Dispatch ────────────────────────────────────────────────────────────────

RENDERERS = {
    "w_000": _render_w000,
    "w_003": _render_w003,
    "w_004": _render_wchart,
    "w_005": _render_w005,
    "w_006": _render_w006,
}


# ─── API pública ─────────────────────────────────────────────────────────────

def generar_word_srs(
    widgets: list[dict],
    nombre_plantilla: str = "Levantamiento de Requerimiento",
    img_raya_b64:   str | None = None,
    img_footer_b64: str | None = None,
) -> bytes:
    doc = Document()

    sec               = doc.sections[0]
    sec.page_width    = Inches(8.5)
    sec.page_height   = Inches(11)
    sec.left_margin   = Inches(1)
    sec.right_margin  = Inches(1)
    sec.top_margin    = Inches(1.2)
    sec.bottom_margin = Inches(1)

    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)

    _add_header(doc, img_raya_b64=img_raya_b64, nombre_plantilla=nombre_plantilla)
    _add_footer(doc, img_footer_b64=img_footer_b64)
    _add_intro(doc)

    for widget in sorted(widgets, key=lambda w: w.get("posicion", 0)):
        renderer = RENDERERS.get(widget.get("id_widget", ""))
        if renderer:
            renderer(doc, widget)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()